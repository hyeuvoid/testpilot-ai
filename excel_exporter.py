from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from models import TestCaseResponse


def build_test_case_excel(result: TestCaseResponse) -> BytesIO:
    """把需求分析和测试用例生成 Excel 文件。"""

    workbook = Workbook()

    # 第一个工作表：需求分析
    analysis_sheet = workbook.active
    analysis_sheet.title = "需求分析"

    analysis_sheet.append(["项目", "内容"])
    analysis_sheet.append(
        ["需求概要", result.requirement_summary]
    )
    analysis_sheet.append(
        [
            "需求点",
            "\n".join(
                f"{item.requirement_id}：{item.description}"
                for item in result.requirement_points
            ),
        ]
    )
    analysis_sheet.append(
        [
            "待确认的歧义",
            "\n".join(
                f"{index}. {item}"
                for index, item in enumerate(
                    result.ambiguities,
                    start=1,
                )
            ),
        ]
    )
    analysis_sheet.append(
        [
            "生成时使用的假设",
            "\n".join(
                f"{index}. {item}"
                for index, item in enumerate(
                    result.assumptions,
                    start=1,
                )
            ),
        ]
    )

    analysis_sheet.column_dimensions["A"].width = 22
    analysis_sheet.column_dimensions["B"].width = 100
    style_sheet(analysis_sheet)

    # 第二个工作表：测试用例
    cases_sheet = workbook.create_sheet("测试用例")

    cases_sheet.append(
        [
            "用例编号",
            "用例标题",
            "关联需求",
            "前置条件",
            "测试步骤及预期结果",
            "优先级",
            "测试类型",
            "基于假设",
        ]
    )

    for test_case in result.test_cases:
        steps_text = "\n".join(
            (
                f"{step.step_number}. "
                f"操作：{step.action}；"
                f"预期：{step.expected_result}"
            )
            for step in test_case.steps
        )

        cases_sheet.append(
            [
                test_case.case_id,
                test_case.title,
                "\n".join(test_case.related_requirement_ids),
                "\n".join(test_case.preconditions),
                steps_text,
                test_case.priority,
                test_case.test_type,
                "是" if test_case.based_on_assumption else "否",
            ]
        )

    column_widths = {
        "A": 20,
        "B": 32,
        "C": 18,
        "D": 42,
        "E": 75,
        "F": 12,
        "G": 16,
        "H": 12,
    }

    for column, width in column_widths.items():
        cases_sheet.column_dimensions[column].width = width

    style_sheet(cases_sheet)

    # 把 Excel 保存到内存中
    excel_file = BytesIO()
    workbook.save(excel_file)
    excel_file.seek(0)

    return excel_file


def style_sheet(sheet) -> None:
    """设置工作表的基础样式。"""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="198754",
    )

    for cell in sheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions