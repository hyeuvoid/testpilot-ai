from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".docx",
    ".pdf",
}

MAX_FILE_SIZE = 5 * 1024 * 1024


class DocumentParseError(ValueError):
    """文档不符合要求或无法解析。"""

    pass


def extract_document_text(
    filename: str,
    file_content: bytes,
) -> str:
    """根据文件类型提取文档中的文本。"""

    file_extension = Path(filename).suffix.lower()

    if file_extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            "只支持 TXT、DOCX 和 PDF 文件"
        )

    if not file_content:
        raise DocumentParseError("上传的文件为空")

    if len(file_content) > MAX_FILE_SIZE:
        raise DocumentParseError(
            "文件不能超过 5MB"
        )

    if file_extension == ".txt":
        text = extract_txt_text(file_content)

    elif file_extension == ".docx":
        text = extract_docx_text(file_content)

    else:
        text = extract_pdf_text(file_content)

    normalized_text = normalize_text(text)

    if len(normalized_text) < 10:
        raise DocumentParseError(
            "没有从文档中提取到足够的文字"
        )

    return normalized_text


def extract_txt_text(file_content: bytes) -> str:
    """读取 TXT 文件，兼容常见中文编码。"""

    encodings = (
        "utf-8-sig",
        "utf-8",
        "gb18030",
    )

    for encoding in encodings:
        try:
            return file_content.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise DocumentParseError(
        "无法识别 TXT 文件的文字编码"
    )


def extract_docx_text(file_content: bytes) -> str:
    """提取 Word 文档中的段落和表格文字。"""

    try:
        document = Document(BytesIO(file_content))
    except Exception as error:
        raise DocumentParseError(
            "Word 文件损坏或无法读取"
        ) from error

    content_parts: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            content_parts.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )

            if row_text:
                content_parts.append(row_text)

    return "\n".join(content_parts)


def extract_pdf_text(file_content: bytes) -> str:
    """提取 PDF 各页面中的文字。"""

    try:
        reader = PdfReader(BytesIO(file_content))

        page_texts = [
            page.extract_text() or ""
            for page in reader.pages
        ]

        return "\n".join(page_texts)

    except Exception as error:
        raise DocumentParseError(
            "PDF 文件损坏、加密或无法读取"
        ) from error


def normalize_text(text: str) -> str:
    """清理空白行和每行两侧的空格。"""

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    return "\n".join(lines)